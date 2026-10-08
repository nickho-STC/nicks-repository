"""
Fill blank MASTER cells in AG:AN, AU:BA, CI and CJ from the clinic visit right
before the CCTA exam date, read (by OCR) from the patient's "clinical note" /
"progress note" files.

Usage (normally run through the .bat files):
    python fill_clinical_notes.py            -> test run, first 10 patients
    python fill_clinical_notes.py --all      -> every patient with blanks
    python fill_clinical_notes.py --reindex  -> rescan the patient folders
    python fill_clinical_notes.py --rows 4249,4250  -> only those Excel rows

Safety rules:
  * the original Excel file is never modified; a new *_NOTES_FILLED.xlsx is written
  * only blank cells are filled (a cell holding just a space counts as blank);
    existing values are never touched
  * values are always one of the column's dropdown choices (Data Dictionary)
  * anything the note does not settle stays blank; every filled value is listed
    in the review log with the note text it came from
"""

import argparse
import csv
import datetime as dt
import difflib
import os
import re
import sys
import time

import cv2

from notes_common import (OUT_DIR, SETTINGS_FILE, FILE_EXTS, TEST_LIMIT, find_files,
                              load_json, load_pages, norm_words, pick_paths,
                              save_json, write_workbook)

HERE = os.path.dirname(os.path.abspath(__file__))
INDEX_FILE = os.path.join(HERE, "note_file_index.json")
OCR_CACHE_FILE = os.path.join(HERE, "note_ocr_cache.json")
TEXT_DIR = os.path.join(OUT_DIR, "visit_text")

# Column -> header the column must have (checked before anything is written)
HEADERS = {
    "AG": "Polysubstance Abuse", "AH": "Autoimmune/Collagen Vascular", "AI": "Vasculitides",
    "AJ": "Thrombo-inflammatory", "AK": "Other Inflammatory Disorder",
    "AL": "Autoinflammatory Syndrome", "AM": "Cardiac Ablation/Date", "AN": "Sleep Apnea",
    "AU": "Pre GLP-1", "AV": "Pre GLP-1 Duration", "AW": "Pre GLP-1 Dose",
    "AX": "Pre PCSK9i/Leqvio", "AY": "Pre PCSK9i Duration", "AZ": "Pre Statin",
    "BA": "Pre Statin Dose", "CI": "SBP (mmHg)", "CJ": "BMI",
}
TARGET_COLS = list(HEADERS)
MAX_COL = 88  # CJ


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def is_blank(v):
    return v is None or str(v).strip() == ""  # strip() also removes non-breaking spaces


# --------------------------------------------------------------------------
# Spreadsheet
# --------------------------------------------------------------------------

def read_targets(excel):
    import openpyxl
    wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
    ws = wb["MASTER"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=MAX_COL))]
    for c, name in HEADERS.items():
        got = str(header[col_index(c)] or "").strip()
        if got != name:
            sys.exit(f"Column {c} header is '{got}', expected '{name}'. "
                     "The spreadsheet layout changed - stopping without changes.")
    targets = []
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=MAX_COL, values_only=True), start=2):
        row = list(row) + [None] * (MAX_COL - len(row))
        if is_blank(row[0]):
            continue
        blanks = [c for c in TARGET_COLS if is_blank(row[col_index(c)])]
        if blanks:
            targets.append({"row": r, "name": str(row[0]).strip(), "exam_date": row[1],
                            "vidistar": row[2], "dob": row[8], "blanks": blanks})
    wb.close()
    return targets


# --------------------------------------------------------------------------
# Find note files
# --------------------------------------------------------------------------

def is_note_file(fn):
    low = fn.lower()
    compact = re.sub(r"[^a-z]", "", low)
    return low.endswith(FILE_EXTS) and ("clinicalnote" in compact or "progressnote" in compact)


def build_index(root):
    print(f"Scanning {root} for clinical/progress note files (only needed once)...")
    files, ndirs, t0 = [], 0, time.time()
    for dirpath, _dirnames, filenames in os.walk(root):
        ndirs += 1
        if ndirs % 2000 == 0:
            print(f"  ...{ndirs} folders scanned, {len(files)} note files found")
        files += [os.path.join(dirpath, fn) for fn in filenames if is_note_file(fn)]
    print(f"  done: {len(files)} note files in {ndirs} folders ({time.time() - t0:.0f}s)")
    save_json(INDEX_FILE, {"root": root, "built": dt.datetime.now().isoformat(timespec="seconds"),
                           "files": files})
    return files


# --------------------------------------------------------------------------
# OCR -> text lines, visit dates
# --------------------------------------------------------------------------

_ocr = None


def ocr(img):
    """OCR boxes. The angle classifier is off: notes are upright, and it flips
    some long printed lines (e.g. the vitals line), which then read as nothing."""
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR(use_angle_cls=False)
    result, _ = _ocr(img)
    tokens = []
    for box, text, _score in result or []:
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        tokens.append({"x0": min(xs), "x1": max(xs), "y0": min(ys), "y1": max(ys), "t": str(text)})
    return tokens


def prep(img):
    if max(img.shape[:2]) > 3000:
        s = 3000 / max(img.shape[:2])
        img = cv2.resize(img, None, fx=s, fy=s, interpolation=cv2.INTER_AREA)
    elif max(img.shape[:2]) < 1600:  # small screenshots: enlarge so OCR can read them
        s = 2200 / max(img.shape[:2])
        img = cv2.resize(img, None, fx=s, fy=s, interpolation=cv2.INTER_CUBIC)
    return img


def to_lines(tokens):
    """Group OCR boxes into text lines. Each line: {"x", "y", "parts": [[x0, text], ...]}."""
    lines = []
    for t in sorted(tokens, key=lambda t: (t["y0"] + t["y1"]) / 2):
        yc, h = (t["y0"] + t["y1"]) / 2, t["y1"] - t["y0"]
        if lines and abs(yc - lines[-1]["yc"]) < 0.5 * max(h, lines[-1]["h"]):
            lines[-1]["toks"].append(t)
        else:
            lines.append({"yc": yc, "h": h, "toks": [t]})
    out = []
    for ln in lines:
        toks = sorted(ln["toks"], key=lambda t: t["x0"])
        out.append({"x": round(toks[0]["x0"]), "y": round(ln["yc"]),
                    "parts": [[round(t["x0"]), t["t"]] for t in toks]})
    return out


MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
DATE_RE = re.compile(
    r"\b(?P<mon>jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s*(?P<d1>\d{1,2})(?:st|nd|rd|th)?,?\s*(?P<y1>\d{4})"
    r"|\b(?P<m2>\d{1,2})\s*[/\-.]\s*(?P<d2>\d{1,2})\s*[/\-.]\s*(?P<y2>\d{4}|\d{2})\b", re.I)
DATE_KEYWORDS = re.compile(r"date\s*of\s*(service|visit)|visit\s*date|service\s*date|encounter\s*date|"
                           r"\bdos\b|appointment\s*date|seen\s*on|date\s*:", re.I)


def find_dates(text):
    """[(date, position, inside_parentheses)] for every date in the text."""
    out = []
    for m in DATE_RE.finditer(text):
        try:
            if m.group("mon"):
                d = dt.date(int(m.group("y1")), MONTHS[m.group("mon").lower()[:3]], int(m.group("d1")))
            else:
                y = int(m.group("y2"))
                if y < 100:
                    y += 2000 if y <= dt.date.today().year % 100 + 1 else 1900
                d = dt.date(y, int(m.group("m2")), int(m.group("d2")))
        except ValueError:
            continue
        before = text[:m.start()]
        inside = before.rfind("(") > before.rfind(")")
        out.append((d, m.start(), inside))
    return out


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        found = find_dates(v)
        return found[0][0] if found else None
    return None


def page_date(lines, dob):
    """Visit date printed on a page: a date next to a 'date of service' style
    keyword, otherwise the first date in the page header that is not the DOB."""
    def ok(d, inside):
        return not inside and d != dob and 2000 <= d.year <= dt.date.today().year + 1
    for ln in lines:
        text = " ".join(p[1] for p in ln["parts"])
        kw = DATE_KEYWORDS.search(text)
        if kw:
            for d, pos, inside in find_dates(text):
                if ok(d, inside) and pos >= kw.start():
                    return d
    for ln in lines:
        text = " ".join(p[1] for p in ln["parts"])
        for d, _pos, inside in find_dates(text):
            if ok(d, inside):
                return d
    return None


def scan_file(path, cache, dob):
    """OCR the top of every page to get each page's visit date (cached)."""
    key = f"{path}|{os.path.getmtime(path) if os.path.exists(path) else 0}"
    if key in cache and "pages" in cache[key]:
        return key
    pages = []
    try:
        for img in load_pages(path):
            img = prep(img)
            lines = to_lines(ocr(img[: int(img.shape[0] * 0.25)]))
            d = page_date(lines, dob)
            pages.append({"date": d.isoformat() if d else None})
    except Exception as e:  # damaged file etc.
        cache[key] = {"error": f"could not read file: {e}", "pages": []}
        save_json(OCR_CACHE_FILE, cache)
        return key
    # pages without a date of their own belong to the visit on the page before
    last = None
    for p in pages:
        p["visit"] = p["date"] or last
        last = p["visit"]
    if pages and pages[0]["visit"] is None:  # nothing on page 1: try the file name
        d = page_date([{"parts": [[0, os.path.basename(path)]]}], dob)
        for p in pages:
            if p["visit"] is None:
                p["visit"] = d.isoformat() if d else None
    cache[key] = {"path": path, "pages": pages}
    save_json(OCR_CACHE_FILE, cache)
    return key


def full_text_lines(path, key, page_numbers, cache):
    """Full OCR of the chosen pages (cached)."""
    full = cache[key].setdefault("full", {})
    todo = [n for n in page_numbers if str(n) not in full]
    if todo:
        for n, img in enumerate(load_pages(path)):
            if n in todo:
                img = prep(img)
                full[str(n)] = {"w": img.shape[1], "lines": to_lines(ocr(img))}
        save_json(OCR_CACHE_FILE, cache)
    return [full[str(n)] for n in page_numbers if str(n) in full]


# --------------------------------------------------------------------------
# Split a visit into sections (Chief Complaint, HPI, PMH, SH, FH, Meds, PE, A/P ...)
# --------------------------------------------------------------------------

SECTIONS = [
    ("signed", r"(full\s*)?progress\s*note\s*signed(\s*by)?|signed\s*by|electronically\s*signed(\s*by)?|"
               r"rendering\s*provider|provider|attending(\s*physician)?|seen\s*by"),
    ("cc", r"chief\s*complaints?|reason\s*for\s*(visit|referral|consult(ation)?)|cc"),
    ("hpi", r"hpi|history\s*of\s*(the\s*)?present\s*illness|interval\s*history|subjective"),
    ("ros", r"ros|review\s*of\s*systems"),
    ("pmh", r"pmh|past\s*medical\s*history|medical\s*history|problem\s*list|active\s*problems|pmhx"),
    ("psh", r"psh|past\s*surgical\s*history|surgical\s*history"),
    ("sh", r"sh|social\s*history|shx"),
    ("fh", r"fh|family\s*history|fhx"),
    ("allergies", r"allergies|allergy|drug\s*allergies"),
    ("meds", r"meds|medications?|current\s*medications?|medication\s*list|outpatient\s*medications?"),
    ("pe", r"pe|physical\s*exam(ination)?|exam|objective|vitals?|vital\s*signs"),
    ("tests", r"test\s*results|results|labs?|diagnostics?|data\s*reviewed"),
    ("ap", r"a\s*/\s*p|a\s*&\s*p|assessment(\s*(and|&)\s*plan)?|impression|diagnos[ie]s|problems\s*addressed"),
    ("goals", r"goals"),
    ("plan", r"plan|recommendations?"),
    ("cpt", r"cpt\s*codes?|procedure\s*codes?"),
]
SECTION_RES = [(k, re.compile(r"^\s*(?:%s)\s*(?::|-|$)\s*(?P<rest>.*)$" % p, re.I)) for k, p in SECTIONS]
NAMES = {"signed": "full progress note signed", "cc": "chief complaint", "hpi": "hpi",
         "ros": "ros", "pmh": "pmh", "psh": "psh", "sh": "sh", "fh": "fh", "allergies": "allergies",
         "meds": "meds", "pe": "pe", "tests": "test results", "ap": "a/p", "goals": "goals",
         "plan": "plan", "cpt": "cptcodes"}
JUNK = re.compile(r"amazing\s*charts|page\s*\d+\s*of\s*\d+|information\s*on\s*this\s*page|"
                  r"confidential|release\s*of\s*this\s*information|written\s*authori[sz]ation|"
                  r"^\s*(item|data|item\s*data)\s*$", re.I)


def heading(text, near_left):
    for key, rx in SECTION_RES:
        m = rx.match(text)
        if m:
            rest = m.group("rest").strip()
            if key == "cc" and re.match(r"(dr|doctor)\b", rest, re.I):
                return None  # "CC: Dr Smith (PCP)" is a copy-to line, not Chief Complaint
            if len(text.strip()) <= 3 and not near_left and not rest:
                return None
            return key, rest
    if near_left and len(text) >= 4:  # OCR-damaged left-column labels ("Chicf Complamt")
        word = re.sub(r"[^a-z/ ]", "", text.lower()).strip()
        best = difflib.get_close_matches(word, list(NAMES.values()), n=1, cutoff=0.8)
        if best:
            return [k for k, v in NAMES.items() if v == best[0]][0], ""
    return None


def split_sections(pages, patient_name):
    last = (patient_name.split(",")[0] or "").strip().lower()
    sections, cur = {}, "top"
    for page in pages:
        lines = page["lines"]
        if not lines:
            continue
        left = sorted(ln["x"] for ln in lines)[len(lines) // 10]
        for ln in lines:
            text = " ".join(p[1] for p in ln["parts"])
            if JUNK.search(text):
                continue
            if ln is lines[0] and last and last in text.lower() and find_dates(text):
                continue  # page header "NAME (DOB)   visit date"
            parts = [p[1] for p in ln["parts"]]
            h = heading(parts[0], ln["x"] <= left + 0.04 * page["w"])
            if h:
                cur, rest = h
                content = " ".join([rest] + parts[1:]).strip()
            else:
                content = text.strip()
            if re.fullmatch(r"by:?", content, re.I):
                continue  # wrapped "Signed / By" label
            sections.setdefault(cur, [])
            if content:
                sections[cur].append(content)
    return {k: "\n".join(v) for k, v in sections.items()}


# --------------------------------------------------------------------------
# Reading facts out of the text
# --------------------------------------------------------------------------

NEG_CUE = re.compile(r"\b(denies|denied|deny|denying|no|not|negative\s*for|without|w/o|never|none|"
                     r"free\s*of|absence\s*of)\b", re.I)
UNSURE_CUE = re.compile(r"\b(evaluate|evaluation|eval|rule\s*out|r/o|exclude|assess|screen(ing)?|"
                        r"suspect(ed)?|possible|probable|concern|risk|probability|predictive|"
                        r"worried|query)\b|\?", re.I)
SCOPE_END = re.compile(r"\b(but|however|although|though|reports|endorses|has|with|positive\s*for)\b", re.I)
AFTER_NEG = re.compile(r"^\s*[:\-]?\s*(denies|denied|no|none|never|negative|neg)\b", re.I)
FAMILY = re.compile(r"\b(family|fhx|father|mother|dad|mom|brother|sister|sibling|son|daughter|"
                    r"uncle|aunt|grand(father|mother|parent))\b", re.I)
DX = ("pmh", "ap", "hpi", "psh", "cc", "top")  # where diagnoses are written
DX_NEED = ("pmh", "ap")  # a visit with one of these lists the patient's conditions


def clauses(text):
    return [c for c in re.split(r"[.;](?=\s|$)|\n", text) if c.strip()]


def mentions(text, term):
    """[('pos'|'neg'|'unsure', clause)] for each clause that mentions term (regex)."""
    out = []
    for cl in clauses(text):
        if FAMILY.search(cl):
            continue  # "mother has lupus" is not the patient's diagnosis
        m = re.search(term, cl, re.I)
        if not m:
            continue
        before = cl[:m.start()]
        neg = list(NEG_CUE.finditer(before))
        if neg and not SCOPE_END.search(before[neg[-1].end():]):
            kind = "neg"
        elif AFTER_NEG.match(cl[m.end():]):
            kind = "neg"
        elif UNSURE_CUE.search(before[-60:]):
            kind = "unsure"
        else:
            kind = "pos"
        out.append((kind, cl.strip()))
    return out


def lookup(sec, term, where=DX):
    """('pos'|'neg'|'unsure'|None, clause) - a positive mention anywhere wins."""
    found = []
    for s in where:
        found += [(k, f"{s.upper()}: {c}") for k, c in mentions(sec.get(s, ""), term)]
    for kind in ("pos", "neg", "unsure"):
        hits = [c for k, c in found if k == kind]
        if hits:
            return kind, hits[0]
    return None, ""


def listed(sec, need):
    present = [s for s in need if sec.get(s, "").strip()]
    return present


def pick_one(sec, options, none_value, where=DX, need=DX_NEED):
    """First option whose term is positively mentioned; none_value when the
    visit lists the patient's conditions and none of them is there."""
    unsure = []
    for value, term in options:
        kind, clause = lookup(sec, term, where)
        if kind == "pos":
            return value, clause
        if kind == "unsure":
            unsure.append(clause)
    if unsure:
        return None, "only mentioned as possible / to evaluate: " + unsure[0]
    present = listed(sec, need)
    if present:
        return none_value, f"not in {'/'.join(s.upper() for s in present)}"
    return None, f"no {'/'.join(s.upper() for s in need)} section in the note"


# --- medications -----------------------------------------------------------

MED_STOPPED = re.compile(r"\b(stopped|discontinued|d/?c'?d|intoleran\w*|allerg\w*|declined|refused|"
                         r"held|holding|not\s+taking|off\s+of|quit|failed)\b", re.I)


def med_lines(sec):
    lines = [ln.strip() for ln in sec.get("meds", "").split("\n") if ln.strip()]
    return [ln for ln in lines if not re.search(r"\bno\s+(current\s+|known\s+)?medications?\b", ln, re.I)]


def find_med(sec, options, none_value):
    """(value, the medication line, next line) for the first listed drug, or none_value."""
    if not sec.get("meds", "").strip():
        return None, "no medication list in the note", ""
    lines = med_lines(sec)
    stopped = []
    for value, term in options:
        for i, ln in enumerate(lines):
            if re.search(term, ln, re.I):
                if MED_STOPPED.search(ln):
                    stopped.append(ln)
                    continue
                return value, ln, lines[i + 1] if i + 1 < len(lines) else ""
    if stopped:  # e.g. "atorvastatin (intolerant)" - not the same as never started
        return None, f"listed as stopped/intolerant - check by hand: {stopped[0]}", ""
    return none_value, "not in medication list", ""


def dose_mg(line, nxt=""):
    for text in (line, nxt):
        m = re.search(r"(\d+(?:\.\d+)?)\s*mg\b", text, re.I)
        if m:
            n = float(m.group(1))
            return f"{int(n) if n == int(n) else n} mg"
    return None


STATINS = [("Rosuvastatin", r"rosuvastatin|crestor|ezallor"), ("Atorvastatin", r"atorvastatin|lipitor|caduet"),
           ("Simvastatin", r"simvastatin|zocor|vytorin"), ("Pravastatin", r"pravastatin|pravachol"),
           ("Lovastatin", r"lovastatin|mevacor|altoprev"), ("Pitavastatin", r"pitavastatin|livalo|zypitamag")]
OTHER_STATIN = r"fluvastatin|lescol"  # not a dropdown choice for Pre Statin
GLP1 = [("Semaglutide", r"semaglutide|ozempic|wegovy|rybelsus"),
        ("Tirzepatide", r"tirzepatide|mounjaro|zepbound"),
        ("Retatrutide", r"retatrutide"),
        ("Dulaglutide/Exenatide", r"dulaglutide|trulicity|exenatide|byetta|bydureon")]
OTHER_GLP1 = r"liraglutide|victoza|saxenda|lixisenatide|adlyxin|soliqua|xultophy"
PCSK9 = [("Repatha", r"repatha|evolocumab"), ("Praluent", r"praluent|alirocumab"),
         ("Leqvio", r"leqvio|inclisiran")]


def duration(drug_term, ctx):
    """Short (<1 year) / Long (>1 year) from the patient's earlier visits:
    on the drug at a visit 1+ year before the CCTA -> Long; not on it at a visit
    within the year before the CCTA (but on it now) -> Short."""
    exam = ctx["exam"]
    if not exam:
        return None, "no exam date"
    year_before = exam - dt.timedelta(days=365)
    on_long_ago, off_recently = None, None
    for d in sorted(ctx["earlier_visits"], reverse=True):
        sec = ctx["visit_sections"](d)
        if not sec.get("meds", "").strip():
            continue
        on = find_med(sec, [(True, drug_term)], False)[0]
        if on and d <= year_before:
            on_long_ago = d
            break
        if not on and d > year_before and off_recently is None:
            off_recently = d
    if on_long_ago:
        return "Long (>1 year)", f"already on it at the {on_long_ago:%m/%d/%Y} visit"
    if off_recently:
        return "Short (<1 year)", f"not on it at the {off_recently:%m/%d/%Y} visit"
    return None, "start date not found in earlier notes - check by hand"


# --- one function per column: (value or None, why) ------------------------

def c_polysubstance(sec, ctx):
    term = (r"polysubstance|substance\s*(ab)?use|substance\s*use\s*disorder|opioid\s*(ab)?use|"
            r"opioid\s*(use\s*disorder|dependence|addiction)|\boud\b|drug\s*(ab)?use|drug\s*addiction|"
            r"heroin|cocaine|crack|methamphetamine|\bmeth\b|fentanyl\s*(ab)?use|\bivdu\b|\bivda\b|"
            r"narcotic\s*(ab)?use|suboxone|buprenorphine|methadone")
    kind, clause = lookup(sec, term, DX + ("sh",))
    if kind == "pos":
        return "Yes", clause
    if kind == "neg":
        return "No", clause
    present = listed(sec, ("pmh", "ap", "sh"))
    if present:
        return "No", f"not in {'/'.join(s.upper() for s in present)}"
    return None, "no PMH/A/P/SH section in the note"


def c_autoimmune(sec, ctx):
    return pick_one(sec, [
        ("Systemic lupus erythematosus (SLE)", r"lupus|\bsle\b"),
        ("Rheumatoid arthritis (RA)", r"rheumatoid"),
        ("Psoriasis / psoriatic arthritis", r"psoria\w*"),
        ("Systemic sclerosis (scleroderma)", r"scleroderma|systemic\s*sclerosis"),
        ("Sjögren’s syndrome", r"sj[oö]gren"),
        ("Mixed connective tissue disease (MCTD)", r"mixed\s*connective|\bmctd\b"),
        ("Dermatomyositis / polymyositis", r"dermatomyositis|polymyositis"),
        ("Primary Raynaud", r"raynaud"),
    ], "None")


def c_vasculitis(sec, ctx):
    return pick_one(sec, [
        ("Granulomatosis with polyangiitis (GPA)", r"granulomatosis\s*with\s*polyangiitis|wegener"),
        ("Eosinophilic granulomatosis with polyangiitis (EGPA)", r"eosinophilic\s*granulomatosis|churg"),
        ("Microscopic polyangiitis (MPA)", r"microscopic\s*polyangiitis"),
        ("Polyarteritis nodosa (PAN)", r"polyarteritis"),
        ("Takayasu arteritis", r"takayasu"),
        ("Giant cell arteritis (GCA)", r"giant\s*cell\s*arteritis|temporal\s*arteritis"),
        ("Behçet disease", r"beh[cç]et"),
        ("Kawasaki disease", r"kawasaki"),
    ], "None")


def c_thrombo(sec, ctx):
    aps, clause = lookup(sec, r"antiphospholipid|anti-?\s*phospholipid|\baps\b|\bapla?\b|lupus\s*anticoagulant")
    if aps == "pos":
        if lookup(sec, r"lupus(?!\s*anticoagulant)|\bsle\b")[0] == "pos":
            return "SLE with antiphospholipid antibodies (aPL-positive SLE)", clause
        return "Antiphospholipid syndrome (APS)", clause
    return pick_one(sec, [], "None")


def c_other_inflam(sec, ctx):
    return pick_one(sec, [
        ("Inflammatory bowel disease (Crohn’s disease / ulcerative colitis)",
         r"crohn|ulcerative\s*colitis|inflammatory\s*bowel|\bibd\b"),
        ("Sarcoidosis", r"sarcoid\w*"),
        ("Ankylosing spondylitis / axial spondyloarthritis", r"ankylosing|spondyloarthr\w*"),
        ("Chronic HIV infection", r"\bhiv\b|human\s*immunodeficiency"),
    ], "None")


def c_autoinflam(sec, ctx):
    return pick_one(sec, [
        ("Familial Mediterranean fever (FMF)", r"mediterranean\s*fever|\bfmf\b"),
        ("TNF receptor–associated periodic syndrome (TRAPS)", r"\btraps\b|tnf\s*receptor"),
        ("Cryopyrin-associated periodic syndromes (CAPS)", r"cryopyrin|\bcaps\b"),
    ], "None")


NOT_CARDIAC = re.compile(r"endometri|uter|vein|venous|varicose|thyroid|liver|hepat|nerve|barrett|"
                         r"esophag|prostat|tumou?r|kidney|renal\s*mass|lung|radiofrequency\s*of\s*the\s*knee", re.I)


def c_ablation(sec, ctx):
    for s in DX:
        for kind, cl in mentions(sec.get(s, ""), r"ablation|\bpvi\b|pulmonary\s*vein\s*isolation|cryoablation"):
            if kind != "pos" or NOT_CARDIAC.search(cl):
                continue
            dates = [d for d, _p, _i in find_dates(cl) if not ctx["exam"] or d <= ctx["exam"]]
            if dates:
                return ("DATE", dates[0]), f"{s.upper()}: {cl}"
            year = re.search(r"\b(19[89]\d|20[0-4]\d)\b", cl)
            if year:
                return year.group(1), f"{s.upper()}: {cl}"
            return "Ablation", f"{s.upper()}: {cl}"
    return pick_one(sec, [], "None", need=("pmh", "psh", "ap"))


def c_apnea(sec, ctx):
    return pick_one(sec, [("Sleep Apnea (yes)", r"sleep\s*apnea|\bosa\b|\bcpap\b|\bbipap\b")], "No")


def c_glp1(sec, ctx):
    value, line, _ = find_med(sec, GLP1, "None")
    if value == "None" and re.search(OTHER_GLP1, "\n".join(med_lines(sec)), re.I):
        return None, "on a GLP-1 that is not a dropdown choice (e.g. liraglutide) - check by hand"
    return value, line


def c_glp1_duration(sec, ctx):
    value, line, _ = find_med(sec, GLP1, "None")
    if value in (None, "None"):
        return value, line
    return duration(dict(GLP1)[value], ctx)


def c_glp1_dose(sec, ctx):
    value, line, _ = find_med(sec, GLP1, "None")
    if value in (None, "None"):
        return value, line
    return None, f"on {value} ({line}) - pick 'Low fixed dose' / 'Titratable dose' by hand"


def c_pcsk9(sec, ctx):
    value, line, _ = find_med(sec, PCSK9, "None")
    return value, line


def c_pcsk9_duration(sec, ctx):
    value, line, _ = find_med(sec, PCSK9, "None")
    if value in (None, "None"):
        return value, line
    return duration(dict(PCSK9)[value], ctx)


def c_statin(sec, ctx):
    value, line, _ = find_med(sec, STATINS, "Not initiated")
    if value == "Not initiated" and re.search(OTHER_STATIN, "\n".join(med_lines(sec)), re.I):
        return None, "on fluvastatin (not a Pre Statin dropdown choice) - check by hand"
    return value, line


def c_statin_dose(sec, ctx):
    value, line, nxt = find_med(sec, STATINS, "Not initiated")
    if value in (None, "Not initiated"):
        return c_statin(sec, ctx)
    d = dose_mg(line, nxt)
    return (d, line) if d else (None, f"dose not readable: {line}")


def c_sbp(sec, ctx):
    v = ctx["vitals"]["sbp"]
    if v is None:
        return None, "no BP in the note"
    if not 90 <= v <= 200:
        return None, f"SBP read as {v} (outside 90-200) - check by hand"
    return v, f"BP {v}/{ctx['vitals']['dbp']}"


def c_bmi(sec, ctx):
    v = ctx["vitals"]["bmi"]
    if v is None:
        return None, "no BMI (or height+weight) in the note"
    if not 12 <= v <= 80:
        return None, f"BMI read as {v} - check by hand"
    return round(float(v), 1), ctx["vitals"]["bmi_from"]


COLUMN_RULES = {
    "AG": c_polysubstance, "AH": c_autoimmune, "AI": c_vasculitis, "AJ": c_thrombo,
    "AK": c_other_inflam, "AL": c_autoinflam, "AM": c_ablation, "AN": c_apnea,
    "AU": c_glp1, "AV": c_glp1_duration, "AW": c_glp1_dose, "AX": c_pcsk9,
    "AY": c_pcsk9_duration, "AZ": c_statin, "BA": c_statin_dose, "CI": c_sbp, "CJ": c_bmi,
}


def num(pattern, text):
    m = re.search(pattern, text, re.I)
    return float(m.group(1)) if m else None


def vitals(sec):
    t = " ".join(sec.get(s, "") for s in ("pe", "top", "hpi", "tests")).replace("\n", " ")
    # OCR often reads l/I/1 for each other: "BMl", "205 Ib"
    v = {"wt": num(r"\b(?:wt|weight)\s*[:.]?\s*(\d{2,3}(?:\.\d)?)\s*(?:[li1]bs?\b|#)", t),
         "ht": num(r"\b(?:ht|height)(?:\s*/\s*\w+)?\s*[:.]?\s*(\d{2}(?:\.\d)?)\s*(?:[il1]n\b|\"|inches)", t),
         "bmi": num(r"\bbm[il1|]\s*[:.]?\s*(\d{2}(?:\.\d+)?)", t)}
    m = re.search(r"\b(?:bp|blood\s*pressure)\s*[:.]?\s*(\d{2,3})\s*[/\\|]\s*(\d{2,3})", t, re.I)
    v["sbp"], v["dbp"] = (int(m.group(1)), int(m.group(2))) if m else (None, None)
    if v["ht"] is None:
        m = re.search(r"\b(?:ht|height)\s*[:.]?\s*(\d)\s*(?:'|’|ft)\s*(\d{1,2})", t, re.I)
        if m:
            v["ht"] = int(m.group(1)) * 12 + int(m.group(2))
    if v["wt"] is None:
        kg = num(r"\b(?:wt|weight)\s*[:.]?\s*(\d{2,3}(?:\.\d)?)\s*kg", t)
        v["wt"] = round(kg * 2.20462, 1) if kg else None
    v["bmi_from"] = f"BMI {v['bmi']}" if v["bmi"] else ""
    if v["bmi"] is None and v["wt"] and v["ht"]:
        v["bmi"] = round(703 * v["wt"] / v["ht"] ** 2, 1)
        v["bmi_from"] = f"calculated from Wt {v['wt']} lb, Ht {v['ht']} in"
    return v


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def find_visits(cands, cache, dob):
    """{date: [(path, cache key, [page numbers])]} for every visit in the files."""
    visits, errors = {}, []
    for path in cands:
        key = scan_file(path, cache, dob)
        if cache[key].get("error"):
            errors.append(cache[key]["error"])
        pages = {}
        for n, p in enumerate(cache[key]["pages"]):
            if p["visit"]:
                pages.setdefault(p["visit"], []).append(n)
        for d, nums in pages.items():
            visits.setdefault(dt.date.fromisoformat(d), []).append((path, key, nums))
    return visits, errors


def pick_visit(dates, exam):
    """The last visit before the exam date (same-day visit only if there is none before)."""
    if not exam:
        return (dates[-1], "no exam date in sheet - used latest visit") if dates else (None, "")
    before = [d for d in dates if d < exam]
    if before:
        return before[-1], ""
    if exam in dates:
        return exam, "visit is ON the exam date (none before)"
    return None, f"no visit before exam date {exam:%m/%d/%Y}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--all", action="store_true", help="process every patient (default: first 10)")
    ap.add_argument("--reindex", action="store_true", help="rescan the patient folders")
    ap.add_argument("--choose", action="store_true", help="choose the Excel file / folder again")
    ap.add_argument("--rows", help="only these Excel rows, e.g. 4249,4250")
    args = ap.parse_args()

    settings = load_json(SETTINGS_FILE, {})
    if args.choose:
        settings = {}
    excel, root = pick_paths(settings)
    os.makedirs(TEXT_DIR, exist_ok=True)

    print(f"Excel file : {excel}")
    print(f"Patients   : {root}")
    targets = read_targets(excel)
    print(f"{len(targets)} patients have blank cells in AG-AN, AU-BA, CI-CJ.")
    if args.rows:
        wanted = {int(r) for r in re.findall(r"\d+", args.rows)}
        targets = [t for t in targets if t["row"] in wanted]
        print(f"Only rows {args.rows}: {len(targets)} patients.")
    elif not args.all:
        targets = targets[:TEST_LIMIT]
        print(f"TEST RUN: only the first {len(targets)} patients.")

    index = load_json(INDEX_FILE, {})
    if args.reindex or index.get("root") != root:
        files = build_index(root)
    else:
        files = index["files"]
        print(f"Using saved list of {len(files)} note files (run with --reindex to rescan).")
    indexed = [(p, set(norm_words(os.path.relpath(p, root))),
                re.sub(r"\D", "", os.path.relpath(p, root))) for p in files]

    cache = load_json(OCR_CACHE_FILE, {})
    fills, log_rows = {}, []
    stats = {"filled": 0, "blank": 0, "no_file": 0, "no_visit": 0}
    for i, t in enumerate(targets, start=1):
        print(f"[{i}/{len(targets)}] row {t['row']}: {t['name']}", end=" ... ", flush=True)
        exam, dob = as_date(t["exam_date"]), as_date(t["dob"])
        log = {"Excel row": t["row"], "Name": t["name"], "VIDISTAR ID": t["vidistar"],
               "Exam date": exam.strftime("%m/%d/%Y") if exam else t["exam_date"]}
        # all of the patient's note files (not only ones named with the exam date)
        cands, how = find_files(dict(t, exam_date=None), indexed)
        log["Matched by"] = how
        if not cands:
            stats["no_file"] += 1
            log["Status"] = "NO CLINICAL/PROGRESS NOTE FILE FOUND"
            print(log["Status"])
            log_rows.append(log)
            continue
        if len(cands) > 60:
            log["Status"] = f"TOO MANY MATCHING FILES ({len(cands)}) - check manually"
            print(log["Status"])
            log_rows.append(log)
            continue

        visits, errors = find_visits(cands, cache, dob)
        dates = sorted(visits)
        visit, note = pick_visit(dates, exam)
        log["Visits found"] = ", ".join(d.strftime("%m/%d/%Y") for d in dates)
        if visit is None:
            stats["no_visit"] += 1
            log["Status"] = (note or "; ".join(errors) or "no visit dates found in the note files").upper()
            print(log["Status"])
            log_rows.append(log)
            continue
        log["Visit used"] = visit.strftime("%m/%d/%Y")
        log["Note file"] = " | ".join(p for p, _, _ in visits[visit])

        sections_cache = {}

        def visit_sections(d):
            if d not in sections_cache:
                pages = []
                for path, key, nums in visits[d]:
                    pages += full_text_lines(path, key, nums, cache)
                sections_cache[d] = split_sections(pages, t["name"])
            return sections_cache[d]

        sec = visit_sections(visit)
        ctx = {"exam": exam, "vitals": vitals(sec), "visit_sections": visit_sections,
               "earlier_visits": [d for d in dates if d < visit]}
        safe = re.sub(r"[^A-Za-z0-9]+", "_", t["name"])[:40]
        with open(os.path.join(TEXT_DIR, f"row{t['row']}_{safe}_{visit:%Y%m%d}.txt"), "w", encoding="utf-8") as f:
            for k, v in sec.items():
                f.write(f"===== {k.upper()} =====\n{v}\n\n")

        row_fill, left_blank = {}, []
        for c in TARGET_COLS:
            col_name = f"{c} {HEADERS[c]}"
            if c not in t["blanks"]:
                log[col_name] = "(already had a value - not changed)"
                continue
            try:
                val, why = COLUMN_RULES[c](sec, ctx)
            except Exception as e:
                val, why = None, f"error: {e}"
            if val is None:
                left_blank.append(c)
                stats["blank"] += 1
                log[col_name] = f"BLANK - {why}"
            else:
                row_fill[c] = val
                stats["filled"] += 1
                shown = val[1].strftime("%m/%d/%Y") if isinstance(val, tuple) else val
                log[col_name] = f"{shown}  <- {why}"
        if row_fill:
            fills[t["row"]] = row_fill
        warn = []
        if note:
            warn.append(note)
        if exam and (exam - visit).days > 60 and "CI" in row_fill:
            warn.append(f"visit is {(exam - visit).days} days before CCTA (SBP wanted within 2 months)")
        if exam and (exam - visit).days > 365 and "CJ" in row_fill:
            warn.append("visit is over 1 year before CCTA (BMI wanted within 1 year)")
        log["Status"] = ("OK" if not left_blank else "PARTIAL: left blank " + ", ".join(left_blank)) + \
            "".join(f" | {w}" for w in warn)
        print(f"visit {visit:%m/%d/%Y}: {log['Status']}")
        log_rows.append(log)

    # --- outputs
    base = os.path.splitext(os.path.basename(excel))[0]
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    tag = "ALL" if args.all and not args.rows else ("ROWS" if args.rows else "TEST")
    out_xlsx = os.path.join(OUT_DIR, f"{base}_NOTES_FILLED_{tag}_{stamp}.xlsx")
    out_log = os.path.join(OUT_DIR, f"notes_review_log_{tag}_{stamp}.csv")
    if fills:
        write_workbook(excel, out_xlsx, fills)
    cols = ["Excel row", "Name", "VIDISTAR ID", "Exam date", "Visit used", "Status", "Visits found",
            "Matched by", "Note file"] + [f"{c} {HEADERS[c]}" for c in TARGET_COLS]
    with open(out_log, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(log_rows)

    print("\n================ SUMMARY ================")
    print(f"Patients processed          : {len(targets)}")
    print(f"Cells filled                : {stats['filled']}")
    print(f"Cells left blank            : {stats['blank']}")
    print(f"Patients with no note file  : {stats['no_file']}")
    print(f"Patients with no visit found: {stats['no_visit']}")
    print(f"Filled Excel (yellow = new) : {out_xlsx if fills else '(none - nothing to fill)'}")
    print(f"Review log                  : {out_log}")
    print(f"Text the script read        : {TEXT_DIR}")
    print("Your original Excel file was NOT changed.")


if __name__ == "__main__":
    main()
